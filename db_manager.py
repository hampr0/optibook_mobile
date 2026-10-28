import os
import json
import pymysql
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from config import (
    DB_HOST, DB_PORT, DB_USER, DB_PASS, DB_NAME,
    BASE_DIR, OFFLINE_QUEUE_FILE, DEFAULT_CHUNK_SIZE
)

DB_FILE = os.path.join(BASE_DIR, "bookst_db.csv")

def safe_int(val, default=0) -> int:
    """콤마, 공백, 소수점이 포함된 문자열을 안전하게 정수로 변환"""
    if val is None:
        return default
    s = str(val).replace(",", "").strip()
    if not s:
        return default
    try:
        return int(float(s))
    except Exception:
        return default

def make_row_key(row_dict: dict) -> str:
    """복본 도서까지 고유하게 식별하기 위한 복합 키 생성"""
    return f"{row_dict.get('거래처','')}|{row_dict.get('납품처','')}|{row_dict.get('등록일시','')}|{row_dict.get('순번','')}|{row_dict.get('ISBN_Clean','')}"

def format_seq_chunk_label(seq_num: int, chunk_size: int = DEFAULT_CHUNK_SIZE) -> str:
    """특정 순번이 속한 1000권 단위 구간 문자열 반환 (예: 1245 -> '1001~2000')"""
    if seq_num <= 0:
        seq_num = 1
    bucket = (seq_num - 1) // chunk_size
    start_no = bucket * chunk_size + 1
    end_no = (bucket + 1) * chunk_size
    return f"{start_no}~{end_no}"

def get_tidb_connection():
    try:
        conn = pymysql.connect(
            host=DB_HOST,
            port=DB_PORT,
            user=DB_USER,
            password=DB_PASS,
            database=DB_NAME,
            charset="utf8mb4",
            ssl={},
            connect_timeout=4,
            read_timeout=6,
            write_timeout=6,
            autocommit=True,
            cursorclass=pymysql.cursors.DictCursor
        )
        return conn
    except Exception:
        return None

def sanitize_df(df_in: pd.DataFrame) -> pd.DataFrame:
    req_cols = [
        '등록일자', '등록일시', '거래처', '납품처', '내용', '순번',
        '도서명', '저자', '출판사', '정가', '권수', 'ISBN',
        'ISBN_Clean', '상태', '입고일시', '비고', 'company_id'
    ]
    if df_in is None or df_in.empty:
        return pd.DataFrame(columns=req_cols)
    for col in req_cols:
        if col not in df_in.columns:
            df_in[col] = ""
    df_in = df_in.fillna("").astype(str)
    mask = (df_in['도서명'].str.strip() != '') & (df_in['상태'].str.strip() != '삭제')
    return df_in[mask].reset_index(drop=True)

class DataManager:
    def __init__(self):
        self.df = pd.DataFrame()
        self.is_online = False
        self.recent_scans = []

    # ==========================================
    # 오프라인 작업 큐 관리 및 자동 동기화
    # ==========================================
    def _load_offline_queue(self) -> list:
        if os.path.exists(OFFLINE_QUEUE_FILE):
            try:
                with open(OFFLINE_QUEUE_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    return data if isinstance(data, list) else []
            except Exception:
                pass
        return []

    def _save_offline_queue(self, queue_list: list):
        try:
            with open(OFFLINE_QUEUE_FILE, "w", encoding="utf-8") as f:
                json.dump(queue_list, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

    def _enqueue_offline_action(self, item: dict):
        q = self._load_offline_queue()
        q.append(item)
        self._save_offline_queue(q)

    def sync_offline_queue(self, conn=None) -> int:
        queue = self._load_offline_queue()
        if not queue:
            return 0

        close_after = False
        if conn is None:
            conn = get_tidb_connection()
            close_after = True
        if not conn:
            return 0

        synced_count = 0
        remaining = []
        try:
            with conn.cursor() as cur:
                for item in queue:
                    try:
                        c = item.get("client", "")
                        d = item.get("delivery", "")
                        rdt = item.get("reg_dt", "")
                        seq = item.get("seq", "")
                        isbn_c = item.get("isbn_clean", "")
                        st = item.get("status", "입고")
                        idt = item.get("intake_dt", "")

                        if seq:
                            cur.execute(
                                "UPDATE books SET status=%s, intake_dt=%s WHERE client=%s AND delivery=%s AND reg_dt=%s AND seq=%s AND isbn_clean=%s",
                                (st, idt, c, d, rdt, seq, isbn_c)
                            )
                        else:
                            cur.execute(
                                "UPDATE books SET status=%s, intake_dt=%s WHERE client=%s AND delivery=%s AND reg_dt=%s AND isbn_clean=%s LIMIT 1",
                                (st, idt, c, d, rdt, isbn_c)
                            )
                        synced_count += 1
                    except Exception:
                        remaining.append(item)
            self._save_offline_queue(remaining)
        except Exception:
            pass
        finally:
            if close_after:
                try:
                    conn.close()
                except Exception:
                    pass
        return synced_count

    def get_pending_sync_count(self) -> int:
        return len(self._load_offline_queue())

    def _save_local_csv(self):
        try:
            self.df.to_csv(DB_FILE, index=False, encoding='utf-8-sig')
        except Exception:
            pass

    def load_database(self, company_id="ALL"):
        conn = get_tidb_connection()
        if conn:
            try:
                self.is_online = True
                self.sync_offline_queue(conn)

                cursor = conn.cursor()
                if company_id == 'ALL':
                    query = """SELECT reg_date AS 등록일자, reg_dt AS 등록일시, client AS 거래처, 
                                      delivery AS 납품처, content AS 내용, seq AS 순번, title AS 도서명, 
                                      author AS 저자, publisher AS 출판사, price AS 정가, qty AS 권수, 
                                      isbn AS ISBN, isbn_clean AS ISBN_Clean, status AS 상태, 
                                      intake_dt AS 입고일시, remarks AS 비고, company_id FROM books"""
                    cursor.execute(query)
                else:
                    query = """SELECT reg_date AS 등록일자, reg_dt AS 등록일시, client AS 거래처, 
                                      delivery AS 납품처, content AS 내용, seq AS 순번, title AS 도서명, 
                                      author AS 저자, publisher AS 출판사, price AS 정가, qty AS 권수, 
                                      isbn AS ISBN, isbn_clean AS ISBN_Clean, status AS 상태, 
                                      intake_dt AS 입고일시, remarks AS 비고, company_id FROM books WHERE company_id=%s"""
                    cursor.execute(query, (company_id,))

                rows = cursor.fetchall()
                conn.close()
                self.df = sanitize_df(pd.DataFrame(rows))
                self._save_local_csv()
                return True
            except Exception:
                try:
                    conn.close()
                except Exception:
                    pass

        self.is_online = False
        if os.path.exists(DB_FILE):
            try:
                local_df = pd.read_csv(DB_FILE, dtype=str, encoding='utf-8-sig')
                self.df = sanitize_df(local_df)
            except Exception:
                self.df = sanitize_df(pd.DataFrame())
        else:
            self.df = sanitize_df(pd.DataFrame())
        return False

    # ==========================================
    # 전체 순번(_eff_seq) 산출 및 1,000권 단위 구간 분석
    # ==========================================
    def _attach_effective_seq(self, df_sub: pd.DataFrame) -> pd.DataFrame:
        """
        데이터프레임에 기준 순번(_eff_seq)을 부여합니다.
        - 도서에 고유 순번(1~N)이 부여되어 있으면 해당 순번을 우선 사용
        - 여러 납품처가 섞여 순번이 중복되거나 순번이 비어 있는 경우 정렬 후 전체 행 순번(1~N) 부여
        """
        if df_sub.empty:
            df_copy = df_sub.copy()
            df_copy['_eff_seq'] = pd.Series(dtype='int64')
            return df_copy

        df_copy = df_sub.copy()
        raw_seqs = [safe_int(x, 0) for x in df_copy['순번']]
        df_copy['_raw_seq'] = raw_seqs

        valid_positive = [x for x in raw_seqs if x > 0]
        unique_ratio = (len(set(valid_positive)) / len(df_copy)) if len(df_copy) > 0 else 0

        # 순번이 고유하게 잘 들어가 있는 경우(단일 납품처 또는 전체 고유 순번)
        if len(valid_positive) == len(df_copy) and unique_ratio >= 0.85:
            df_copy['_eff_seq'] = df_copy['_raw_seq']
            df_copy.sort_values(by=['_eff_seq'], ascending=True, inplace=True)
        else:
            # 복수 납품처 통합 조회이거나 순번 누락 행이 있을 때: 원본/순번 정렬 후 1~N 전체 순번 부여
            df_copy['_sort_seq'] = [x if x > 0 else 9999999 for x in raw_seqs]
            df_copy.sort_values(by=['납품처', '_sort_seq'], ascending=[True, True], inplace=True)
            df_copy['_eff_seq'] = range(1, len(df_copy) + 1)

        return df_copy

    def get_range_options(self, vendor: str = "전체", chunk_size: int = DEFAULT_CHUNK_SIZE) -> list:
        """드롭다운용 1,000권 단위 구간 옵션 리스트 반환 (예: ['전체 구간', '1~1000', '1001~2000', ...])"""
        if self.df.empty:
            return ["전체 구간"]
        sub = self.df if vendor == "전체" else self.df[self.df['납품처'] == vendor]
        if sub.empty:
            return ["전체 구간"]

        sub_seq = self._attach_effective_seq(sub)
        max_seq = int(sub_seq['_eff_seq'].max()) if not sub_seq.empty else 0
        if max_seq <= 0:
            return ["전체 구간"]

        total_buckets = (max_seq - 1) // chunk_size + 1
        options = ["전체 구간"]
        for b in range(total_buckets):
            s_no = b * chunk_size + 1
            e_no = (b + 1) * chunk_size
            options.append(f"{s_no}~{e_no}")
        return options

    def get_chunk_summary(self, vendor: str = "전체", chunk_size: int = DEFAULT_CHUNK_SIZE) -> list:
        """
        전체 또는 특정 납품처 도서를 1,000권 단위(1~1000, 1001~2000...)로 나누어
        각 구간의 전체 권수, 입고 권수, 미입고 권수, 입고율(%), 금액 합계를 반환
        """
        if self.df.empty:
            return []
        sub = self.df if vendor == "전체" else self.df[self.df['납품처'] == vendor]
        if sub.empty:
            return []

        sub_seq = self._attach_effective_seq(sub)
        sub_seq['_bucket'] = (sub_seq['_eff_seq'] - 1) // chunk_size
        sub_seq['_price_int'] = [safe_int(x, 0) for x in sub_seq['정가']]

        chunks = []
        for bucket_idx, grp in sub_seq.groupby('_bucket', sort=True):
            b_idx = int(bucket_idx)
            start_no = b_idx * chunk_size + 1
            end_no = (b_idx + 1) * chunk_size
            range_key = f"{start_no}~{end_no}"
            label = f"{start_no:,} ~ {end_no:,}번"

            t_cnt = len(grp)
            i_cnt = len(grp[grp['상태'] == '입고'])
            p_cnt = t_cnt - i_cnt
            rate_float = round((i_cnt / t_cnt * 100.0), 1) if t_cnt > 0 else 0.0
            rate_int = int(i_cnt / t_cnt * 100) if t_cnt > 0 else 0
            tot_price = int(grp['_price_int'].sum())

            chunks.append({
                "range_key": range_key,
                "label": label,
                "start": start_no,
                "end": end_no,
                "total": t_cnt,
                "intake": i_cnt,
                "pending": p_cnt,
                "rate": rate_int,
                "rate_float": rate_float,
                "price": tot_price
            })
        return chunks

    def _filter_by_seq_range(self, df_with_seq: pd.DataFrame, seq_range: str) -> pd.DataFrame:
        if not seq_range or seq_range == "전체 구간" or "~" not in seq_range:
            return df_with_seq
        try:
            parts = seq_range.replace(",", "").split("~")
            start_no = int(parts[0].strip())
            end_no = int(parts[1].strip())
            return df_with_seq[(df_with_seq['_eff_seq'] >= start_no) & (df_with_seq['_eff_seq'] <= end_no)]
        except Exception:
            return df_with_seq

    # ==========================================
    # DB 업데이트 및 바코드 스캔 처리
    # ==========================================
    def _update_row_status_in_db(self, target_idx: int, new_status: str, now_dt: str) -> dict:
        self.df.at[target_idx, '상태'] = new_status
        self.df.at[target_idx, '입고일시'] = now_dt
        row_data = self.df.loc[target_idx].to_dict()
        self._save_local_csv()

        c = str(row_data.get('거래처', '')).strip()
        d = str(row_data.get('납품처', '')).strip()
        rdt = str(row_data.get('등록일시', '')).strip()
        seq = str(row_data.get('순번', '')).strip()
        isbn_clean = str(row_data.get('ISBN_Clean', '')).strip()

        conn = get_tidb_connection()
        if conn:
            try:
                self.sync_offline_queue(conn)
                with conn.cursor() as cur:
                    if seq:
                        cur.execute(
                            "UPDATE books SET status=%s, intake_dt=%s WHERE client=%s AND delivery=%s AND reg_dt=%s AND seq=%s AND isbn_clean=%s",
                            (new_status, now_dt, c, d, rdt, seq, isbn_clean)
                        )
                    else:
                        cur.execute(
                            "UPDATE books SET status=%s, intake_dt=%s WHERE client=%s AND delivery=%s AND reg_dt=%s AND isbn_clean=%s LIMIT 1",
                            (new_status, now_dt, c, d, rdt, isbn_clean)
                        )
                conn.close()
                self.is_online = True
            except Exception:
                self.is_online = False
                self._enqueue_offline_action({
                    "client": c, "delivery": d, "reg_dt": rdt, "seq": seq,
                    "isbn_clean": isbn_clean, "status": new_status, "intake_dt": now_dt
                })
        else:
            self.is_online = False
            self._enqueue_offline_action({
                "client": c, "delivery": d, "reg_dt": rdt, "seq": seq,
                "isbn_clean": isbn_clean, "status": new_status, "intake_dt": now_dt
            })

        return row_data

    def process_barcode_scan(self, isbn_input: str, target_vendor: str = "전체", seq_range: str = "전체 구간"):
        isbn_clean = ''.join(filter(str.isdigit, str(isbn_input)))
        if not isbn_clean or self.df.empty:
            return "NOT_FOUND", None, self.is_online

        vendor_pool = self.df if target_vendor == "전체" else self.df[self.df['납품처'] == target_vendor]
        if vendor_pool.empty:
            return "NOT_FOUND", None, self.is_online

        vendor_pool_seq = self._attach_effective_seq(vendor_pool)
        search_pool = self._filter_by_seq_range(vendor_pool_seq, seq_range)

        # 1. 정확히 일치하는 ISBN 검색
        matches = search_pool[search_pool['ISBN_Clean'] == isbn_clean]

        # 2. 4~6자리 짧은 숫자를 입력한 경우 (끝자리 일치 검색)
        if matches.empty and 4 <= len(isbn_clean) <= 6:
            suffix_matches = search_pool[search_pool['ISBN_Clean'].str.endswith(isbn_clean)]
            pending_suffix = suffix_matches[suffix_matches['상태'] != '입고']
            if len(pending_suffix) == 1:
                matches = pending_suffix
            elif len(pending_suffix) > 1:
                if pending_suffix['ISBN_Clean'].nunique() == 1:
                    matches = pending_suffix
                else:
                    return "MULTIPLE_SUFFIX", {"count": len(pending_suffix), "query": isbn_clean}, self.is_online
            elif len(suffix_matches) >= 1:
                matches = suffix_matches

        # 현재 선택 구간에 없으면 다른 구간 또는 다른 납품처에 있는지 확인
        if matches.empty:
            in_other_range = vendor_pool_seq[vendor_pool_seq['ISBN_Clean'] == isbn_clean]
            if in_other_range.empty and 4 <= len(isbn_clean) <= 6:
                in_other_range = vendor_pool_seq[vendor_pool_seq['ISBN_Clean'].str.endswith(isbn_clean)]
            if not in_other_range.empty and seq_range != "전체 구간":
                other_seq = int(in_other_range.iloc[0].get('_eff_seq', 1))
                other_chunk = format_seq_chunk_label(other_seq)
                return "WRONG_RANGE", {"target_range": other_chunk, "seq": other_seq}, self.is_online

            other_matches = self.df[self.df['ISBN_Clean'] == isbn_clean]
            if other_matches.empty and 4 <= len(isbn_clean) <= 6:
                other_matches = self.df[self.df['ISBN_Clean'].str.endswith(isbn_clean)]
            if not other_matches.empty:
                other_vendor = other_matches.iloc[0].get('납품처', '')
                return "WRONG_VENDOR", {"target": other_vendor}, self.is_online
            return "NOT_FOUND", None, self.is_online

        pending_matches = matches[matches['상태'] != '입고']
        if pending_matches.empty:
            already_idx = matches.index[0]
            already_row = self.df.loc[already_idx].to_dict()
            eff_s = int(matches.loc[already_idx, '_eff_seq'])
            already_row['_eff_seq'] = eff_s
            already_row['_chunk_label'] = format_seq_chunk_label(eff_s)
            return "ALREADY_IN", already_row, self.is_online

        target_idx = pending_matches.index[0]
        eff_seq_val = int(pending_matches.loc[target_idx, '_eff_seq'])
        now_dt = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        row_data = self._update_row_status_in_db(target_idx, '입고', now_dt)

        row_data['_remaining_copies'] = len(pending_matches) - 1
        row_data['_eff_seq'] = eff_seq_val
        row_data['_chunk_label'] = format_seq_chunk_label(eff_seq_val)

        self.recent_scans = [r for r in self.recent_scans if make_row_key(r) != make_row_key(row_data)]
        self.recent_scans.insert(0, row_data)
        if len(self.recent_scans) > 30:
            self.recent_scans.pop()

        return "SUCCESS", row_data, self.is_online

    def toggle_book_intake_by_key(self, row_key: str) -> tuple:
        if self.df.empty:
            return False, None, self.is_online

        parts = row_key.split("|")
        if len(parts) != 5:
            return False, None, self.is_online

        c, d, rdt, seq, isbn_c = parts
        mask = (
            (self.df['거래처'] == c) &
            (self.df['납품처'] == d) &
            (self.df['등록일시'] == rdt) &
            (self.df['순번'] == seq) &
            (self.df['ISBN_Clean'] == isbn_c)
        )
        matched = self.df[mask]
        if matched.empty:
            return False, None, self.is_online

        target_idx = matched.index[0]
        current_status = self.df.at[target_idx, '상태']
        if current_status == '입고':
            new_status = '미입고'
            now_dt = ""
        else:
            new_status = '입고'
            now_dt = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        updated_row = self._update_row_status_in_db(target_idx, new_status, now_dt)

        if new_status == '미입고':
            self.recent_scans = [r for r in self.recent_scans if make_row_key(r) != row_key]
        else:
            self.recent_scans = [r for r in self.recent_scans if make_row_key(r) != row_key]
            self.recent_scans.insert(0, updated_row)

        return True, updated_row, self.is_online

    def get_vendor_progress(self, vendor: str = "전체", seq_range: str = "전체 구간") -> dict:
        """선택된 납품처 및 순번 구간의 실시간 검수 진행률 반환"""
        if self.df.empty:
            return {"total": 0, "intake": 0, "pending": 0, "rate": 0, "rate_float": 0.0}
        sub = self.df if vendor == "전체" else self.df[self.df['납품처'] == vendor]
        if sub.empty:
            return {"total": 0, "intake": 0, "pending": 0, "rate": 0, "rate_float": 0.0}

        if seq_range != "전체 구간":
            sub = self._attach_effective_seq(sub)
            sub = self._filter_by_seq_range(sub, seq_range)

        t = len(sub)
        i = len(sub[sub['상태'] == '입고'])
        p = t - i
        r_int = int(i / t * 100) if t > 0 else 0
        r_flt = round(i / t * 100.0, 1) if t > 0 else 0.0
        return {"total": t, "intake": i, "pending": p, "rate": r_int, "rate_float": r_flt}

    def get_summary_data(self):
        if self.df.empty:
            return {"total_qty": 0, "intake_qty": 0, "pending_qty": 0, "total_price": 0, "vendors": []}

        t_all = len(self.df)
        d_all = len(self.df[self.df['상태'] == '입고'])
        p_all = t_all - d_all
        price_all = sum(safe_int(x) for x in self.df['정가'])

        unique_batches = self.df[['거래처', '납품처']].drop_duplicates().reset_index(drop=True)
        vendors = []
        for _, row in unique_batches.iterrows():
            c, v = str(row['거래처']), str(row['납품처'])
            mask = (self.df['거래처'] == c) & (self.df['납품처'] == v)
            v_df = self.df[mask]
            v_t = len(v_df)
            v_d = len(v_df[v_df['상태'] == '입고'])
            v_p = v_t - v_d
            v_pr = sum(safe_int(x) for x in v_df['정가'])

            vendors.append({
                "client": c,
                "delivery": v,
                "total": v_t,
                "intake": v_d,
                "pending": v_p,
                "price": v_pr
            })

        return {
            "total_qty": t_all,
            "intake_qty": d_all,
            "pending_qty": p_all,
            "total_price": price_all,
            "vendors": vendors
        }

    def get_filtered_books(self, vendor="전체", status="전체", search_query="", seq_range="전체 구간"):
        if self.df.empty:
            return []
        f_df = self.df if vendor == "전체" else self.df[self.df['납품처'] == vendor]
        if f_df.empty:
            return []

        f_df = self._attach_effective_seq(f_df)
        if seq_range != "전체 구간":
            f_df = self._filter_by_seq_range(f_df, seq_range)

        if status != "전체":
            f_df = f_df[f_df['상태'] == status]
        if search_query:
            q = search_query.strip().lower()
            mask = (
                f_df['도서명'].astype(str).str.lower().str.contains(q, na=False) |
                f_df['ISBN_Clean'].astype(str).str.contains(q, na=False) |
                f_df['저자'].astype(str).str.lower().str.contains(q, na=False) |
                f_df['순번'].astype(str).str.contains(q, na=False)
            )
            f_df = f_df[mask]
        return f_df.to_dict('records')

    def get_history_books(self, target_date: str = ""):
        if self.df.empty:
            return []
        if not target_date:
            target_date = datetime.now().strftime("%Y-%m-%d")

        mask = (self.df['상태'] == '입고') & (self.df['입고일시'].astype(str).str.startswith(target_date))
        h_df = self.df[mask].copy()
        if h_df.empty:
            return []
        h_df.sort_values(by='입고일시', ascending=False, inplace=True)
        return h_df.to_dict('records')

    def get_vendor_list(self):
        if self.df.empty or '납품처' not in self.df.columns:
            return ["전체"]
        vendors = [v for v in self.df['납품처'].dropna().unique() if str(v).strip() and str(v).strip() != '납품처']
        return ["전체"] + sorted(vendors)

    def export_books_to_excel(self, book_list: list, vendor_name: str = "전체") -> str:
        if not book_list:
            return ""

        def sort_key(item):
            if '_eff_seq' in item:
                return safe_int(item.get('_eff_seq'), default=999999)
            return safe_int(item.get('순번', 0), default=999999)

        sorted_books = sorted(book_list, key=sort_key)

        export_data = []
        for idx, row in enumerate(sorted_books, start=1):
            seq_val = safe_int(row.get('순번', ''), default=safe_int(row.get('_eff_seq', idx), idx))
            price_val = safe_int(row.get('정가', 0), default=0)
            qty_val = safe_int(row.get('권수', 1), default=1)

            isbn_raw = str(row.get('ISBN_Clean', '')).strip()
            if isbn_raw.isdigit() and len(isbn_raw) <= 15:
                isbn_val = int(isbn_raw)
            else:
                isbn_val = isbn_raw

            export_data.append({
                '전번(전체번호)': idx,
                '거래처': row.get('거래처', ''),
                '납품처': row.get('납품처', ''),
                '내용': row.get('내용', ''),
                '순번': seq_val,
                '도서명': row.get('도서명', ''),
                '저자': row.get('저자', ''),
                '출판사': row.get('출판사', ''),
                '정가': price_val,
                '권수': qty_val,
                'ISBN': isbn_val,
                '입고상태': row.get('상태', ''),
                '작업일자': row.get('입고일시', '')
            })

        export_df = pd.DataFrame(export_data)
        time_tag = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_vendor = "".join(c for c in str(vendor_name) if c.isalnum() or c in (' ', '_', '-', '~')).strip() or "전체"
        export_filename = f"입고검수내역_{safe_vendor}_{time_tag}.xlsx"

        candidate_dirs = [
            "/storage/emulated/0/Download",
            os.environ.get("FLET_APP_STORAGE_DATA", ""),
            BASE_DIR
        ]

        save_path = os.path.join(BASE_DIR, export_filename)
        for d in candidate_dirs:
            if d and os.path.exists(d):
                test_path = os.path.join(d, export_filename)
                try:
                    self._write_excel_file(export_df, test_path)
                    return test_path
                except Exception:
                    continue

        self._write_excel_file(export_df, save_path)
        return save_path

    def _write_excel_file(self, export_df: pd.DataFrame, save_path: str):
        with pd.ExcelWriter(save_path, engine='openpyxl') as writer:
            export_df.to_excel(writer, index=False, sheet_name='입고검수내역')
            ws = writer.sheets['입고검수내역']

            col_widths = {
                'A': 6, 'B': 12, 'C': 14, 'D': 10, 'E': 6,
                'F': 32, 'G': 18, 'H': 16, 'I': 11, 'J': 6,
                'K': 15, 'L': 9, 'M': 18
            }
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width

            thin_border = Border(
                left=Side(style='thin', color='A0A0A0'),
                right=Side(style='thin', color='A0A0A0'),
                top=Side(style='thin', color='A0A0A0'),
                bottom=Side(style='thin', color='A0A0A0')
            )
            header_fill = PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid')
            header_font = Font(name='맑은 고딕', size=10, bold=True, color='1565C0')
            body_font = Font(name='맑은 고딕', size=10)

            ws.row_dimensions[1].height = 25
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
                cell.border = thin_border

            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                ws.row_dimensions[row[0].row].height = 20
                for col_idx, cell in enumerate(row, start=1):
                    cell.font = body_font
                    cell.border = thin_border
                    col_name = export_df.columns[col_idx - 1]

                    if col_name in ['도서명', '저자', '출판사', '내용']:
                        cell.alignment = Alignment(horizontal='left', vertical='center', shrink_to_fit=True)
                    elif col_name in ['정가', '권수']:
                        cell.alignment = Alignment(horizontal='right', vertical='center')
                        if col_name == '정가' and isinstance(cell.value, (int, float)):
                            cell.number_format = '#,##0'
                    elif col_name == 'ISBN':
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = '0'
                    elif col_name == '입고상태':
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        if str(cell.value) == '입고':
                            cell.font = Font(name='맑은 고딕', size=10, bold=True, color='2E7D32')
                        else:
                            cell.font = Font(name='맑은 고딕', size=10, bold=True, color='C62828')
                    else:
                        cell.alignment = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)

db_service = DataManager()